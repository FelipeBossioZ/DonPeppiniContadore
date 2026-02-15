# 🎩 Don Peppini Contadore - Medios Magnéticos DIAN
# backend/contabilidad/medios_magneticos.py
#
# Formatos: 1001, 1003, 1005, 1006, 1007, 1008, 1009, 1012, 2276
# Resolución 000124 de 2021 y modificaciones vigentes.

from decimal import Decimal
from io import BytesIO
from collections import defaultdict

from django.db.models import Sum, Q
from django.db.models.functions import Coalesce
from django.http import HttpResponse

from rest_framework import views
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from empresas.models import Empresa
from terceros.models import Tercero
from contabilidad.models import MovimientoContable, AsientoContable


# =============================================================================
# HELPERS
# =============================================================================

def get_tercero_mm_data(tercero):
    """Extrae datos de un Tercero en formato Medios Magnéticos DIAN."""
    if not tercero:
        return {
            'tipo_documento': '', 'numero_documento': '', 'dv': '',
            'primer_apellido': '', 'segundo_apellido': '',
            'primer_nombre': '', 'otros_nombres': '',
            'razon_social': 'SIN TERCERO',
            'direccion': '', 'codigo_dpto': '', 'codigo_mcp': '',
            'codigo_pais': '169',
        }

    es_juridica = tercero.tipo_documento == 'NIT'

    # Tipo doc DIAN: 13=CC, 31=NIT, 22=CE, 41=PA, 11=RC, 12=TI, 21=TE, 42=DIE
    TIPO_DOC_DIAN = {
        'CC': '13', 'NIT': '31', 'CE': '22', 'PA': '41',
        'TI': '12', 'RC': '11', 'TE': '21', 'DIE': '42',
    }

    return {
        'tipo_documento': TIPO_DOC_DIAN.get(tercero.tipo_documento, '13'),
        'numero_documento': tercero.numero_documento or '',
        'dv': tercero.digito_verificacion or '',
        'primer_apellido': '' if es_juridica else (tercero.primer_apellido or ''),
        'segundo_apellido': '' if es_juridica else (tercero.segundo_apellido or ''),
        'primer_nombre': '' if es_juridica else (tercero.primer_nombre or ''),
        'otros_nombres': '' if es_juridica else (tercero.otros_nombres or ''),
        'razon_social': tercero.nombre_razon_social if es_juridica else '',
        'direccion': tercero.direccion or '',
        'codigo_dpto': tercero.codigo_departamento or '',
        'codigo_mcp': tercero.codigo_municipio or '',
        'codigo_pais': tercero.codigo_pais or '169',
    }


def resolver_tercero(mov):
    """Devuelve el tercero del movimiento (línea) o el del asiento como fallback."""
    return mov.tercero or mov.asiento.tercero


# =============================================================================
# MAPEO PUC → CONCEPTOS DIAN (Formato 1001)
# =============================================================================

_CONCEPTO_1001_RULES = [
    # --- Aportes patronales (6 dígitos, antes de genéricos 4-dígito) ---
    ('510568', '5031'),  # ARL
    ('510569', '5030'),  # Salud empleador
    ('510570', '5027'),  # Pensión empleador
    ('510572', '5032'),  # Caja compensación
    ('510575', '5033'),  # ICBF
    ('510578', '5034'),  # SENA
    ('520568', '5031'),
    ('520569', '5030'),
    ('520570', '5027'),
    ('520572', '5032'),
    ('520575', '5033'),
    ('520578', '5034'),
    # --- Provisiones prestaciones sociales → 5001 ---
    ('510530', '5001'), ('510533', '5001'), ('510536', '5001'), ('510539', '5001'),
    ('520530', '5001'), ('520533', '5001'), ('520536', '5001'), ('520539', '5001'),
    # --- Gastos de personal genéricos → 5001 ---
    ('5105', '5001'),
    ('5205', '5001'),
    # --- Honorarios ---
    ('5110', '5002'), ('5210', '5002'),
    # --- Impuestos ---
    ('5115', '5012'), ('5215', '5012'),
    # --- Arrendamientos ---
    ('5120', '5005'), ('5220', '5005'),
    # --- Contribuciones y afiliaciones ---
    ('5125', '5023'), ('5225', '5023'),
    # --- Seguros ---
    ('5130', '5011'), ('5230', '5011'),
    # --- Servicios ---
    ('5135', '5004'), ('5235', '5004'),
    # --- Gastos legales ---
    ('5140', '5004'), ('5240', '5004'),
    # --- Mantenimiento ---
    ('5145', '5004'), ('5245', '5004'),
    # --- Adecuaciones ---
    ('5150', '5004'), ('5250', '5004'),
    # --- Gastos de viaje ---
    ('5155', '5013'), ('5255', '5013'),
    # --- Depreciaciones y amortizaciones: NO generan pago real ---
    ('5160', None), ('5165', None), ('5260', None), ('5265', None),
    # --- Financieros ---
    ('5170', '5006'), ('5270', '5006'),
    # --- Diversos ---
    ('5195', '5013'), ('5295', '5013'),
    # --- Costos → Compras ---
    ('6', '5016'),
    # --- Fallback clase 5 ---
    ('5', '5013'),
]


def _concepto_1001(codigo_cuenta):
    cod = str(codigo_cuenta)
    for prefix, concepto in _CONCEPTO_1001_RULES:
        if cod.startswith(prefix):
            return concepto
    return '5013'


_CONCEPTO_1007_RULES = [
    ('4175', '4002'),   # Devoluciones en ventas
    ('42', '4002'),     # Ingresos no operacionales
    ('41', '4001'),     # Ingresos operacionales
    ('4', '4001'),      # Fallback
]


def _concepto_1007(codigo_cuenta):
    cod = str(codigo_cuenta)
    for prefix, concepto in _CONCEPTO_1007_RULES:
        if cod.startswith(prefix):
            return concepto
    return '4001'


# =============================================================================
# ESTILOS EXCEL
# =============================================================================

_HEADER_FONT = Font(bold=True, color='FFFFFF', size=10)
_HEADER_FILL = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
_HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
_BORDER_THIN = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
_MONEY_FMT = '#,##0'


def _aplicar_encabezado(ws, row_num, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _BORDER_THIN


def _titulo_hoja(ws, empresa, titulo, year):
    ws.cell(row=1, column=1, value=empresa.razon_social).font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=titulo).font = Font(bold=True, size=11, color='4F46E5')
    ws.cell(row=3, column=1, value=f'Año Gravable {year}')


def _escribir_headers(ws, headers, fila=4):
    for col, h in enumerate(headers, 1):
        ws.cell(row=fila, column=col, value=h)
    _aplicar_encabezado(ws, fila, len(headers))
    return fila + 1


def _sin_datos(ws, row, fila_inicio):
    if row == fila_inicio:
        ws.cell(row=fila_inicio, column=1, value='SIN DATOS PARA ESTE FORMATO')


def _fmt_money(ws, row, cols):
    for c in cols:
        ws.cell(row=row, column=c).number_format = _MONEY_FMT


# =============================================================================
# QUERIES BASE
# =============================================================================

def _movimientos_periodo(empresa, fecha_inicio, fecha_fin, cuenta_filter):
    return MovimientoContable.objects.filter(
        asiento__empresa=empresa,
        asiento__fecha__gte=fecha_inicio,
        asiento__fecha__lte=fecha_fin,
        asiento__estado='vigente',
        **cuenta_filter
    ).select_related('asiento', 'asiento__tercero', 'tercero', 'cuenta')


def _saldos_corte(empresa, fecha_corte, cuenta_regex):
    return MovimientoContable.objects.filter(
        asiento__empresa=empresa,
        asiento__fecha__lte=fecha_corte,
        asiento__estado='vigente',
        cuenta__codigo__regex=cuenta_regex
    ).select_related('asiento', 'asiento__tercero', 'tercero', 'cuenta')


# =============================================================================
# FORMATO 1001 — Pagos o abonos en cuenta y retenciones practicadas
# =============================================================================

def generar_1001(wb, empresa, fecha_inicio, fecha_fin, year):
    ws = wb.create_sheet('1001')
    _titulo_hoja(ws, empresa, '1001 - Pagos o abonos en cuenta y retenciones en la fuente practicadas', year)

    headers = [
        'Concepto', 'Tipo documento', 'Número identificación',
        'Primer apellido', 'Segundo apellido', 'Primer nombre', 'Otros nombres',
        'Razón social', 'Dirección', 'Cód dpto', 'Cód mpio',
        'País residencia', 'Pago deducible', 'Pago NO deducible',
        'IVA mayor valor costo deducible', 'IVA mayor valor costo NO deducible',
        'Retención Renta practicada', 'Retención Renta asumida',
        'Retención IVA régimen común', 'Retención IVA no domiciliados'
    ]
    fila = _escribir_headers(ws, headers)
    money_cols = list(range(13, 21))

    # 1) Pagos: débitos en cuentas 5xxx y 6xxx
    movs = _movimientos_periodo(empresa, fecha_inicio, fecha_fin, {'cuenta__codigo__regex': r'^[56]'})
    datos = {}
    for mov in movs:
        concepto = _concepto_1001(mov.cuenta.codigo)
        if concepto is None:
            continue
        tercero = resolver_tercero(mov)
        if not tercero:
            continue
        key = (tercero.id, concepto)
        if key not in datos:
            datos[key] = {'tercero': tercero, 'concepto': concepto,
                          'pago': Decimal('0'), 'rete_renta': Decimal('0'), 'rete_iva': Decimal('0')}
        datos[key]['pago'] += mov.debito

    # 2) Retenciones Renta practicadas (2365xx)
    for mov in _movimientos_periodo(empresa, fecha_inicio, fecha_fin, {'cuenta__codigo__startswith': '2365'}):
        tercero = resolver_tercero(mov)
        if not tercero:
            continue
        matched = False
        for key in datos:
            if key[0] == tercero.id:
                datos[key]['rete_renta'] += mov.credito
                matched = True
                break
        if not matched:
            key = (tercero.id, '5013')
            if key not in datos:
                datos[key] = {'tercero': tercero, 'concepto': '5013',
                              'pago': Decimal('0'), 'rete_renta': Decimal('0'), 'rete_iva': Decimal('0')}
            datos[key]['rete_renta'] += mov.credito

    # 3) Retenciones IVA practicadas (2367xx)
    for mov in _movimientos_periodo(empresa, fecha_inicio, fecha_fin, {'cuenta__codigo__startswith': '2367'}):
        tercero = resolver_tercero(mov)
        if not tercero:
            continue
        for key in datos:
            if key[0] == tercero.id:
                datos[key]['rete_iva'] += mov.credito
                break

    row = fila
    for d in sorted(datos.values(), key=lambda x: x['concepto']):
        mm = get_tercero_mm_data(d['tercero'])
        ws.cell(row=row, column=1, value=d['concepto'])
        ws.cell(row=row, column=2, value=mm['tipo_documento'])
        ws.cell(row=row, column=3, value=mm['numero_documento'])
        ws.cell(row=row, column=4, value=mm['primer_apellido'])
        ws.cell(row=row, column=5, value=mm['segundo_apellido'])
        ws.cell(row=row, column=6, value=mm['primer_nombre'])
        ws.cell(row=row, column=7, value=mm['otros_nombres'])
        ws.cell(row=row, column=8, value=mm['razon_social'])
        ws.cell(row=row, column=9, value=mm['direccion'])
        ws.cell(row=row, column=10, value=mm['codigo_dpto'])
        ws.cell(row=row, column=11, value=mm['codigo_mcp'])
        ws.cell(row=row, column=12, value=mm['codigo_pais'])
        ws.cell(row=row, column=13, value=float(d['pago']))
        ws.cell(row=row, column=14, value=0)
        ws.cell(row=row, column=15, value=0)
        ws.cell(row=row, column=16, value=0)
        ws.cell(row=row, column=17, value=float(d['rete_renta']))
        ws.cell(row=row, column=18, value=0)
        ws.cell(row=row, column=19, value=float(d['rete_iva']))
        ws.cell(row=row, column=20, value=0)
        _fmt_money(ws, row, money_cols)
        row += 1

    _sin_datos(ws, row, fila)
    return row - fila


# =============================================================================
# FORMATO 1003 — Retenciones que le practicaron
# =============================================================================

def generar_1003(wb, empresa, fecha_inicio, fecha_fin, year):
    ws = wb.create_sheet('1003')
    _titulo_hoja(ws, empresa, '1003 - Retenciones en la fuente que le practicaron', year)

    headers = [
        'Concepto', 'Tipo documento', 'Número identificación', 'DV',
        'Primer apellido', 'Segundo apellido', 'Primer nombre', 'Otros nombres',
        'Razón social', 'Dirección', 'Cód dpto', 'Cód mpio',
        'Pago sujeto a retención', 'Retención que le practicaron'
    ]
    fila = _escribir_headers(ws, headers)

    # 1355xx = Anticipo impuestos (retenciones que nos practicaron)
    movs = _movimientos_periodo(empresa, fecha_inicio, fecha_fin, {'cuenta__codigo__startswith': '1355'})
    datos = {}
    for mov in movs:
        tercero = resolver_tercero(mov)
        if not tercero:
            continue
        key = tercero.id
        if key not in datos:
            datos[key] = {'tercero': tercero, 'base': Decimal('0'), 'retencion': Decimal('0')}
        datos[key]['retencion'] += mov.debito

    # Base: ingresos (4xxx créditos) por tercero
    ingresos = _movimientos_periodo(empresa, fecha_inicio, fecha_fin, {
        'cuenta__codigo__startswith': '4', 'credito__gt': 0
    })
    bases = defaultdict(Decimal)
    for mov in ingresos:
        tercero = resolver_tercero(mov)
        if tercero:
            bases[tercero.id] += mov.credito
    for key, d in datos.items():
        d['base'] = bases.get(key, Decimal('0'))

    row = fila
    for d in datos.values():
        mm = get_tercero_mm_data(d['tercero'])
        ws.cell(row=row, column=1, value='1003')
        ws.cell(row=row, column=2, value=mm['tipo_documento'])
        ws.cell(row=row, column=3, value=mm['numero_documento'])
        ws.cell(row=row, column=4, value=mm['dv'])
        ws.cell(row=row, column=5, value=mm['primer_apellido'])
        ws.cell(row=row, column=6, value=mm['segundo_apellido'])
        ws.cell(row=row, column=7, value=mm['primer_nombre'])
        ws.cell(row=row, column=8, value=mm['otros_nombres'])
        ws.cell(row=row, column=9, value=mm['razon_social'])
        ws.cell(row=row, column=10, value=mm['direccion'])
        ws.cell(row=row, column=11, value=mm['codigo_dpto'])
        ws.cell(row=row, column=12, value=mm['codigo_mcp'])
        ws.cell(row=row, column=13, value=float(d['base']))
        ws.cell(row=row, column=14, value=float(d['retencion']))
        _fmt_money(ws, row, [13, 14])
        row += 1

    _sin_datos(ws, row, fila)
    return row - fila


# =============================================================================
# FORMATOS 1005/1006 — IVA Descontable / Generado
# =============================================================================

def _generar_iva(wb, empresa, fecha_inicio, fecha_fin, year, sheet, titulo, es_compra):
    ws = wb.create_sheet(sheet)
    _titulo_hoja(ws, empresa, titulo, year)

    headers = [
        'Tipo documento', 'Número identificación', 'DV',
        'Primer apellido', 'Segundo apellido', 'Primer nombre', 'Otros nombres',
        'Razón social',
        'IVA Descontable' if es_compra else 'IVA Generado',
        'IVA devoluciones', 'IVA mayor valor costo' if es_compra else 'Impuesto al consumo'
    ]
    fila = _escribir_headers(ws, headers)

    filtro = {'cuenta__codigo__startswith': '2408'}
    if es_compra:
        filtro['debito__gt'] = 0
    else:
        filtro['credito__gt'] = 0

    movs = _movimientos_periodo(empresa, fecha_inicio, fecha_fin, filtro)
    datos = {}
    for mov in movs:
        tercero = resolver_tercero(mov)
        if not tercero:
            continue
        key = tercero.id
        if key not in datos:
            datos[key] = {'tercero': tercero, 'iva': Decimal('0')}
        datos[key]['iva'] += mov.debito if es_compra else mov.credito

    row = fila
    for d in datos.values():
        mm = get_tercero_mm_data(d['tercero'])
        ws.cell(row=row, column=1, value=mm['tipo_documento'])
        ws.cell(row=row, column=2, value=mm['numero_documento'])
        ws.cell(row=row, column=3, value=mm['dv'])
        ws.cell(row=row, column=4, value=mm['primer_apellido'])
        ws.cell(row=row, column=5, value=mm['segundo_apellido'])
        ws.cell(row=row, column=6, value=mm['primer_nombre'])
        ws.cell(row=row, column=7, value=mm['otros_nombres'])
        ws.cell(row=row, column=8, value=mm['razon_social'])
        ws.cell(row=row, column=9, value=float(d['iva']))
        ws.cell(row=row, column=10, value=0)
        ws.cell(row=row, column=11, value=0)
        _fmt_money(ws, row, [9, 10, 11])
        row += 1

    _sin_datos(ws, row, fila)
    return row - fila


def generar_1005(wb, empresa, fecha_inicio, fecha_fin, year):
    return _generar_iva(wb, empresa, fecha_inicio, fecha_fin, year,
                        '1005', '1005 - IVA Descontable', es_compra=True)


def generar_1006(wb, empresa, fecha_inicio, fecha_fin, year):
    return _generar_iva(wb, empresa, fecha_inicio, fecha_fin, year,
                        '1006', '1006 - IVA Generado', es_compra=False)


# =============================================================================
# FORMATO 1007 — Ingresos Recibidos
# =============================================================================

def generar_1007(wb, empresa, fecha_inicio, fecha_fin, year):
    ws = wb.create_sheet('1007')
    _titulo_hoja(ws, empresa, '1007 - Ingresos recibidos', year)

    headers = [
        'Concepto', 'Tipo documento', 'Número identificación',
        'Primer apellido', 'Segundo apellido', 'Primer nombre', 'Otros nombres',
        'Razón social', 'País residencia',
        'Ingresos brutos', 'Devoluciones rebajas descuentos'
    ]
    fila = _escribir_headers(ws, headers)

    movs = _movimientos_periodo(empresa, fecha_inicio, fecha_fin, {
        'cuenta__codigo__startswith': '4', 'credito__gt': 0
    })
    datos = {}
    for mov in movs:
        tercero = resolver_tercero(mov)
        if not tercero:
            continue
        concepto = _concepto_1007(mov.cuenta.codigo)
        key = (tercero.id, concepto)
        if key not in datos:
            datos[key] = {'tercero': tercero, 'concepto': concepto,
                          'ingresos': Decimal('0'), 'devoluciones': Decimal('0')}
        datos[key]['ingresos'] += mov.credito

    # Devoluciones: débitos en 4175xx
    for mov in _movimientos_periodo(empresa, fecha_inicio, fecha_fin, {
        'cuenta__codigo__startswith': '4175', 'debito__gt': 0
    }):
        tercero = resolver_tercero(mov)
        if not tercero:
            continue
        for key in datos:
            if key[0] == tercero.id:
                datos[key]['devoluciones'] += mov.debito
                break

    row = fila
    for d in sorted(datos.values(), key=lambda x: x['concepto']):
        mm = get_tercero_mm_data(d['tercero'])
        ws.cell(row=row, column=1, value=d['concepto'])
        ws.cell(row=row, column=2, value=mm['tipo_documento'])
        ws.cell(row=row, column=3, value=mm['numero_documento'])
        ws.cell(row=row, column=4, value=mm['primer_apellido'])
        ws.cell(row=row, column=5, value=mm['segundo_apellido'])
        ws.cell(row=row, column=6, value=mm['primer_nombre'])
        ws.cell(row=row, column=7, value=mm['otros_nombres'])
        ws.cell(row=row, column=8, value=mm['razon_social'])
        ws.cell(row=row, column=9, value=mm['codigo_pais'])
        ws.cell(row=row, column=10, value=float(d['ingresos']))
        ws.cell(row=row, column=11, value=float(d['devoluciones']))
        _fmt_money(ws, row, [10, 11])
        row += 1

    _sin_datos(ws, row, fila)
    return row - fila


# =============================================================================
# FORMATOS 1008, 1009, 1012 — Saldos a corte
# =============================================================================

def _generar_saldos(wb, empresa, fecha_corte, year, sheet, titulo, cuenta_regex,
                    es_deudor, concepto_default, tiene_ubicacion):
    ws = wb.create_sheet(sheet)
    _titulo_hoja(ws, empresa, titulo, year)

    headers = ['Concepto', 'Tipo documento', 'Número identificación', 'DV',
               'Primer apellido', 'Segundo apellido', 'Primer nombre', 'Otros nombres',
               'Razón social']
    if tiene_ubicacion:
        headers += ['Dirección', 'Cód dpto', 'Cód mpio']
    headers += ['País residencia', f'Saldo al 31-12']
    fila = _escribir_headers(ws, headers)
    col_saldo = len(headers)

    movs = _saldos_corte(empresa, fecha_corte, cuenta_regex)
    datos = {}
    for mov in movs:
        tercero = resolver_tercero(mov)
        if not tercero:
            continue
        key = tercero.id
        if key not in datos:
            datos[key] = {'tercero': tercero, 'deb': Decimal('0'), 'cre': Decimal('0'),
                          'cuenta': mov.cuenta.codigo}
        datos[key]['deb'] += mov.debito
        datos[key]['cre'] += mov.credito

    row = fila
    for d in datos.values():
        saldo = (d['deb'] - d['cre']) if es_deudor else (d['cre'] - d['deb'])
        if saldo <= 0:
            continue

        mm = get_tercero_mm_data(d['tercero'])

        concepto = concepto_default
        if sheet == '1012':
            cod = d['cuenta']
            if cod.startswith('1110'):
                concepto = '1110'
            elif cod.startswith('1120'):
                concepto = '1120'
            elif cod.startswith('12'):
                concepto = '1200'

        col = 1
        ws.cell(row=row, column=col, value=concepto); col += 1
        ws.cell(row=row, column=col, value=mm['tipo_documento']); col += 1
        ws.cell(row=row, column=col, value=mm['numero_documento']); col += 1
        ws.cell(row=row, column=col, value=mm['dv']); col += 1
        ws.cell(row=row, column=col, value=mm['primer_apellido']); col += 1
        ws.cell(row=row, column=col, value=mm['segundo_apellido']); col += 1
        ws.cell(row=row, column=col, value=mm['primer_nombre']); col += 1
        ws.cell(row=row, column=col, value=mm['otros_nombres']); col += 1
        ws.cell(row=row, column=col, value=mm['razon_social']); col += 1
        if tiene_ubicacion:
            ws.cell(row=row, column=col, value=mm['direccion']); col += 1
            ws.cell(row=row, column=col, value=mm['codigo_dpto']); col += 1
            ws.cell(row=row, column=col, value=mm['codigo_mcp']); col += 1
        ws.cell(row=row, column=col, value=mm['codigo_pais']); col += 1
        ws.cell(row=row, column=col, value=float(saldo))
        ws.cell(row=row, column=col).number_format = _MONEY_FMT
        row += 1

    _sin_datos(ws, row, fila)
    return row - fila


def generar_1008(wb, empresa, fecha_corte, year):
    return _generar_saldos(wb, empresa, fecha_corte, year,
                           '1008', '1008 - Saldo cuentas por cobrar', r'^13',
                           es_deudor=True, concepto_default='1315', tiene_ubicacion=True)


def generar_1009(wb, empresa, fecha_corte, year):
    return _generar_saldos(wb, empresa, fecha_corte, year,
                           '1009', '1009 - Saldo cuentas por pagar', r'^2[23]',
                           es_deudor=False, concepto_default='2205', tiene_ubicacion=True)


def generar_1012(wb, empresa, fecha_corte, year):
    return _generar_saldos(wb, empresa, fecha_corte, year,
                           '1012', '1012 - Inversiones, títulos y cuentas', r'^1[12]',
                           es_deudor=True, concepto_default='1110', tiene_ubicacion=False)


# =============================================================================
# FORMATO 2276 — Rentas de Trabajo (desde LiquidacionEmpleado)
# =============================================================================

def generar_2276(wb, empresa, fecha_inicio, fecha_fin, year):
    from nomina.models import Empleado, LiquidacionEmpleado, Nomina

    ws = wb.create_sheet('2276')
    _titulo_hoja(ws, empresa, '2276 - Información de ingresos y retenciones por rentas de trabajo', year)

    headers = [
        'NIT Informante', 'Tipo doc beneficiario', 'Nro identificación',
        'Primer apellido', 'Segundo apellido', 'Primer nombre', 'Otros nombres',
        'Dirección', 'Cód dpto', 'Cód mpio', 'País',
        'Pagos por salarios',           # 12
        'Pagos por emolumentos',        # 13
        'Pagos con bonos',              # 14
        'Exceso alimentación 41 UVT',   # 15
        'Pagos por honorarios',         # 16
        'Pagos por servicios',          # 17
        'Pagos por comisiones',         # 18
        'Pagos prestaciones sociales',  # 19
        'Pagos por viáticos',           # 20
        'Pagos gastos representación',  # 21
        'Compensaciones cooperativo',   # 22
        'Apoyos económicos',            # 23
        'Otros pagos',                  # 24
        'Cesantías pagadas',            # 25
        'Cesantías consignadas fondo',  # 26
        'Cesantías régimen tradicional',# 27
        'Pensiones',                    # 28
        'Total ingresos brutos',        # 29
        'Aportes salud obligatoria',    # 30
        'Aportes pensión obligatoria',  # 31
        'Aportes voluntarios RAIS',     # 32
        'Aportes voluntarios pensión',  # 33
        'Aportes AFC',                  # 34
        'Aportes AVC',                  # 35
        'Retención renta',              # 36
    ]
    fila = _escribir_headers(ws, headers)
    money_cols = list(range(12, 37))

    nominas = Nomina.objects.filter(empresa=empresa, anio=year)
    liquidaciones = LiquidacionEmpleado.objects.filter(
        nomina__in=nominas
    ).select_related('empleado', 'empleado__tercero')

    acumulado = {}
    for liq in liquidaciones:
        eid = liq.empleado.id
        if eid not in acumulado:
            acumulado[eid] = {
                'tercero': liq.empleado.tercero,
                'salarios': Decimal('0'), 'comisiones': Decimal('0'),
                'prestaciones': Decimal('0'), 'otros': Decimal('0'),
                'cesantias': Decimal('0'), 'total_bruto': Decimal('0'),
                'salud': Decimal('0'), 'pension': Decimal('0'),
                'retencion': Decimal('0'),
            }
        a = acumulado[eid]
        a['salarios'] += liq.salario_devengado + liq.auxilio_transporte + liq.horas_extras + liq.recargos
        a['comisiones'] += liq.comisiones
        a['prestaciones'] += liq.provision_prima + liq.provision_vacaciones
        a['otros'] += liq.bonificaciones + liq.otros_devengados
        a['cesantias'] += liq.provision_cesantias + liq.provision_int_cesantias
        a['total_bruto'] += liq.total_devengado
        a['salud'] += liq.salud_empleado
        a['pension'] += liq.pension_empleado + liq.fsp
        a['retencion'] += liq.retencion_fuente

    row = fila
    for a in acumulado.values():
        if a['total_bruto'] <= 0:
            continue
        mm = get_tercero_mm_data(a['tercero'])
        ws.cell(row=row, column=1, value=empresa.nit)
        ws.cell(row=row, column=2, value=mm['tipo_documento'])
        ws.cell(row=row, column=3, value=mm['numero_documento'])
        ws.cell(row=row, column=4, value=mm['primer_apellido'])
        ws.cell(row=row, column=5, value=mm['segundo_apellido'])
        ws.cell(row=row, column=6, value=mm['primer_nombre'])
        ws.cell(row=row, column=7, value=mm['otros_nombres'])
        ws.cell(row=row, column=8, value=mm['direccion'])
        ws.cell(row=row, column=9, value=mm['codigo_dpto'])
        ws.cell(row=row, column=10, value=mm['codigo_mcp'])
        ws.cell(row=row, column=11, value=mm['codigo_pais'])
        ws.cell(row=row, column=12, value=float(a['salarios']))
        for c in range(13, 16):
            ws.cell(row=row, column=c, value=0)
        ws.cell(row=row, column=16, value=0)
        ws.cell(row=row, column=17, value=0)
        ws.cell(row=row, column=18, value=float(a['comisiones']))
        ws.cell(row=row, column=19, value=float(a['prestaciones']))
        for c in range(20, 24):
            ws.cell(row=row, column=c, value=0)
        ws.cell(row=row, column=24, value=float(a['otros']))
        ws.cell(row=row, column=25, value=0)
        ws.cell(row=row, column=26, value=float(a['cesantias']))
        ws.cell(row=row, column=27, value=0)
        ws.cell(row=row, column=28, value=0)
        ws.cell(row=row, column=29, value=float(a['total_bruto']))
        ws.cell(row=row, column=30, value=float(a['salud']))
        ws.cell(row=row, column=31, value=float(a['pension']))
        for c in range(32, 36):
            ws.cell(row=row, column=c, value=0)
        ws.cell(row=row, column=36, value=float(a['retencion']))
        _fmt_money(ws, row, money_cols)
        row += 1

    _sin_datos(ws, row, fila)
    return row - fila


# =============================================================================
# VISTA PRINCIPAL
# =============================================================================

FORMATOS_INFO = {
    '1001': {'nombre': 'Pagos y retenciones practicadas', 'version': 10},
    '1003': {'nombre': 'Retenciones que le practicaron', 'version': 7},
    '1005': {'nombre': 'IVA Descontable', 'version': 8},
    '1006': {'nombre': 'IVA Generado', 'version': 8},
    '1007': {'nombre': 'Ingresos Recibidos', 'version': 9},
    '1008': {'nombre': 'Cuentas por Cobrar', 'version': 7},
    '1009': {'nombre': 'Cuentas por Pagar', 'version': 7},
    '1012': {'nombre': 'Inversiones y Cuentas', 'version': 7},
    '2276': {'nombre': 'Rentas de Trabajo', 'version': 2},
}


class MediosMagneticosView(views.APIView):
    """
    🎩 Generación de Medios Magnéticos DIAN

    GET ?empresa=X&year=Y          → Excel con TODOS los formatos
    GET ?empresa=X&year=Y&formato=1001 → JSON preview del formato
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        empresa_id = request.query_params.get('empresa')
        year = request.query_params.get('year')
        formato = request.query_params.get('formato')

        if not empresa_id or not year:
            return Response({'error': 'Empresa y año son requeridos'}, status=400)

        try:
            empresa = Empresa.objects.get(id=empresa_id)
            year = int(year)
        except Empresa.DoesNotExist:
            return Response({'error': 'Empresa no encontrada'}, status=404)
        except ValueError:
            return Response({'error': 'Año inválido'}, status=400)

        fecha_inicio = f'{year}-01-01'
        fecha_fin = f'{year}-12-31'

        if formato:
            return self._preview(empresa, fecha_inicio, fecha_fin, year, formato)
        return self._excel(empresa, fecha_inicio, fecha_fin, year)

    def _preview(self, empresa, fecha_inicio, fecha_fin, year, formato):
        if formato not in FORMATOS_INFO:
            return Response({'error': f'Formato {formato} no soportado'}, status=400)

        wb = Workbook()
        wb.remove(wb.active)
        self._run(wb, empresa, fecha_inicio, fecha_fin, year, formato)

        ws = wb[formato]
        columnas = []
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row=4, column=col).value
            if val:
                columnas.append(str(val))

        registros = []
        for r in range(5, ws.max_row + 1):
            fila = {}
            vacia = True
            for i, cn in enumerate(columnas):
                val = ws.cell(row=r, column=i + 1).value
                if val is not None and str(val) != 'SIN DATOS PARA ESTE FORMATO':
                    vacia = False
                fila[cn] = val
            if not vacia:
                registros.append(fila)

        return Response({
            'formato': formato,
            'nombre': FORMATOS_INFO[formato]['nombre'],
            'version': FORMATOS_INFO[formato]['version'],
            'year': year,
            'total_registros': len(registros),
            'columnas': columnas,
            'registros': registros,
        })

    def _excel(self, empresa, fecha_inicio, fecha_fin, year):
        wb = Workbook()
        wb.remove(wb.active)
        for fmt in FORMATOS_INFO:
            self._run(wb, empresa, fecha_inicio, fecha_fin, year, fmt)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        resp = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        resp['Content-Disposition'] = f'attachment; filename=MediosMagneticos_{empresa.nit}_{year}.xlsx'
        return resp

    def _run(self, wb, empresa, fi, ff, year, fmt):
        fc = ff  # fecha_corte = fecha_fin para saldos
        dispatch = {
            '1001': lambda: generar_1001(wb, empresa, fi, ff, year),
            '1003': lambda: generar_1003(wb, empresa, fi, ff, year),
            '1005': lambda: generar_1005(wb, empresa, fi, ff, year),
            '1006': lambda: generar_1006(wb, empresa, fi, ff, year),
            '1007': lambda: generar_1007(wb, empresa, fi, ff, year),
            '1008': lambda: generar_1008(wb, empresa, fc, year),
            '1009': lambda: generar_1009(wb, empresa, fc, year),
            '1012': lambda: generar_1012(wb, empresa, fc, year),
            '2276': lambda: generar_2276(wb, empresa, fi, ff, year),
        }
        return dispatch.get(fmt, lambda: 0)()
