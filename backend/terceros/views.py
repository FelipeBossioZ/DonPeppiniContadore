# 🎩 Don Peppini Contadore - Views de Terceros (GLOBALES)
from rest_framework import viewsets, filters, status, views
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.decorators import action
from django.db.models import Q
from .models import Tercero, EmpresaTercero
from .serializers import (
    TerceroListSerializer,
    TerceroDetailSerializer,
    TerceroCreateSerializer,
    TerceroSerializer,
)


class TerceroViewSet(viewsets.ModelViewSet):
    """
    🎩 ViewSet para gestión de Terceros (GLOBALES)

    Los terceros son compartidos entre todas las empresas.
    """
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['nombre_razon_social', 'numero_documento', 'primer_apellido', 'primer_nombre']
    ordering_fields = ['nombre_razon_social', 'numero_documento', 'tipo_tercero']
    ordering = ['nombre_razon_social']

    def get_serializer_class(self):
        if self.action == 'create':
            return TerceroCreateSerializer
        return TerceroDetailSerializer

    def get_queryset(self):
        base = Tercero.objects.all() if self.request.query_params.get('incluir_inactivos') == 'true' else Tercero.objects.filter(activo=True)

        # Filtrar por tipo de tercero (opcional)
        tipo = self.request.query_params.get('tipo') or self.request.query_params.get('tipo_tercero')
        if tipo:
            base = base.filter(tipo_tercero=tipo)

        # Filtrar por empresa si se proporciona
        empresa_id = self.request.query_params.get('empresa')
        if empresa_id:
            base = base.filter(
                Q(es_compartido=True) | Q(empresas_rel__empresa_id=empresa_id)
            ).distinct()

        return base

    def create(self, request, *args, **kwargs):
        """Crear tercero y vincularlo a la empresa actual si se especifica"""
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        tercero = serializer.save()
        
        # Vincular a la empresa actual (si se envió empresa_id)
        empresa_id = request.data.get('empresa_id') or request.query_params.get('empresa')
        if empresa_id:
            from empresas.models import Empresa
            try:
                emp = Empresa.objects.get(pk=empresa_id)
                EmpresaTercero.objects.get_or_create(empresa=emp, tercero=tercero)
            except Exception:
                pass  # Si falla la empresa, igual se creó el tercero
        
        # Retornar con serializer completo para que el frontend tenga todos los campos
        return Response(
            TerceroDetailSerializer(tercero).data,
            status=status.HTTP_201_CREATED
        )

    @action(detail=True, methods=['post'])
    def desactivar(self, request, pk=None):
        """Desactivar un tercero (soft delete)"""
        tercero = self.get_object()
        tercero.activo = False
        tercero.save()
        return Response({'status': 'Tercero desactivado'})

    @action(detail=True, methods=['post'])
    def activar(self, request, pk=None):
        """Activar un tercero"""
        tercero = self.get_object()
        tercero.activo = True
        tercero.save()
        return Response({'status': 'Tercero activado'})

    @action(detail=False, methods=['get'])
    def buscar(self, request):
        """
        Búsqueda rápida: GET /api/terceros/buscar/?q=banco
        """
        q = request.query_params.get('q', '').strip()
        if len(q) < 2:
            return Response([])

        qs = Tercero.objects.filter(
            Q(numero_documento__icontains=q) |
            Q(nombre_razon_social__icontains=q) |
            Q(primer_nombre__icontains=q) |
            Q(primer_apellido__icontains=q)
        ).filter(activo=True)[:20]

        return Response(TerceroListSerializer(qs, many=True).data)

    @action(detail=False, methods=['get'])
    def exportar(self, request):
        """Exportar terceros de una empresa a Excel."""
        from empresas.models import Empresa
        from django.http import HttpResponse
        from io import BytesIO
        import openpyxl
        from openpyxl.styles import Font

        empresa_id = request.query_params.get('empresa')
        if not empresa_id:
            return Response({'error': 'empresa requerido'}, status=400)

        try:
            emp = Empresa.objects.get(pk=empresa_id)
        except Empresa.DoesNotExist:
            return Response({'error': 'Empresa no encontrada'}, status=404)

        terceros = Tercero.objects.filter(
            empresas_rel__empresa=emp
        ).order_by('nombre_razon_social')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f'Terceros {emp.razon_social[:20]}'

        headers = [
            'tipo_documento', 'numero_documento', 'digito_verificacion',
            'nombre_razon_social', 'primer_apellido', 'segundo_apellido',
            'primer_nombre', 'otros_nombres', 'tipo_tercero',
            'direccion', 'ciudad', 'departamento',
            'codigo_municipio', 'codigo_departamento', 'codigo_pais',
            'telefono', 'email',
            'es_autoretenedor', 'es_gran_contribuyente',
            'es_declarante', 'regimen_simple',
        ]
        ws.append(headers)

        bool_to_yn = lambda v: 'Si' if v else 'No'
        for t in terceros:
            ws.append([
                t.tipo_documento or '',
                t.numero_documento or '',
                t.digito_verificacion or '',
                t.nombre_razon_social or '',
                t.primer_apellido or '',
                t.segundo_apellido or '',
                t.primer_nombre or '',
                t.otros_nombres or '',
                t.tipo_tercero or '',
                t.direccion or '',
                t.ciudad or '',
                t.departamento or '',
                t.codigo_municipio or '',
                t.codigo_departamento or '',
                t.codigo_pais or '',
                t.telefono or '',
                t.email or '',
                bool_to_yn(t.es_autoretenedor),
                bool_to_yn(t.es_gran_contribuyente),
                bool_to_yn(t.es_declarante),
                bool_to_yn(t.regimen_simple),
            ])

        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.font = Font(bold=True)

        # Data validation: tipo_tercero dropdown (columna I = 9)
        from openpyxl.worksheet.datavalidation import DataValidation
        tipos_tercero = '"CLI,PRO,EMP,SOC,OTR"'
        dv = DataValidation(type='list', formula1=tipos_tercero, allow_blank=True)
        dv.error = 'Seleccione: CLI, PRO, EMP, SOC u OTR'
        dv.errorTitle = 'Tipo invalido'
        dv.prompt = 'Seleccione tipo de tercero'
        dv.promptTitle = 'Tipo de Tercero'
        last_row = ws.max_row
        if last_row > 1:
            dv.add(f'I2:I{last_row}')
        ws.add_data_validation(dv)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = (
            f'attachment; filename="terceros_{emp.razon_social[:20]}.xlsx"'
        )
        return response

    @action(detail=False, methods=['post'])
    def importar(self, request):
        """Importar terceros desde Excel a una empresa."""
        from empresas.models import Empresa

        empresa_id = request.data.get('empresa')
        archivo = request.FILES.get('archivo')

        if not empresa_id or not archivo:
            return Response({'error': 'empresa y archivo requeridos'}, status=400)

        try:
            emp = Empresa.objects.get(pk=empresa_id)
        except Empresa.DoesNotExist:
            return Response({'error': 'Empresa no encontrada'}, status=404)

        try:
            import pandas as pd
            df = pd.read_excel(archivo, sheet_name=0)
        except Exception as e:
            return Response({'error': f'Error leyendo Excel: {str(e)}'}, status=400)

        creados = 0
        vinculados = 0
        omitidos = 0
        errores = []

        def safe_str(v):
            """Return string or empty."""
            return str(v).strip() if pd.notna(v) and v != '' else ''

        def is_true(v):
            raw = str(v).strip().lower() if pd.notna(v) else ''
            return raw in ('1', 'si', 'yes', 'true', 'verdadero')

        for idx, row in df.iterrows():
            if pd.isna(row.get('nombre_razon_social')) or pd.isna(row.get('numero_documento')):
                continue

            num_doc = safe_str(row['numero_documento']).replace('.0', '')
            nombre = safe_str(row['nombre_razon_social'])

            # If tercero already exists globally, just link to this empresa
            existente = Tercero.objects.filter(numero_documento=num_doc).first()
            if existente:
                _, created = EmpresaTercero.objects.get_or_create(
                    empresa=emp, tercero=existente
                )
                if created:
                    vinculados += 1
                else:
                    omitidos += 1
                continue

            try:
                tercero = Tercero(
                    tipo_documento=safe_str(row.get('tipo_documento')) or 'NIT',
                    numero_documento=num_doc,
                    digito_verificacion=safe_str(row.get('digito_verificacion')),
                    nombre_razon_social=nombre,
                    primer_apellido=safe_str(row.get('primer_apellido')),
                    segundo_apellido=safe_str(row.get('segundo_apellido')),
                    primer_nombre=safe_str(row.get('primer_nombre')),
                    otros_nombres=safe_str(row.get('otros_nombres')),
                    tipo_tercero=safe_str(row.get('tipo_tercero')) or 'OTR',
                    direccion=safe_str(row.get('direccion')),
                    ciudad=safe_str(row.get('ciudad')),
                    departamento=safe_str(row.get('departamento')),
                    codigo_municipio=safe_str(row.get('codigo_municipio')),
                    codigo_departamento=safe_str(row.get('codigo_departamento')),
                    codigo_pais=safe_str(row.get('codigo_pais')) or '169',
                    telefono=safe_str(row.get('telefono')),
                    email=safe_str(row.get('email')),
                    es_autoretenedor=is_true(row.get('es_autoretenedor')),
                    es_gran_contribuyente=is_true(row.get('es_gran_contribuyente')),
                    es_declarante=is_true(row.get('es_declarante')),
                    regimen_simple=is_true(row.get('regimen_simple')),
                )
                tercero.save()
                EmpresaTercero.objects.create(empresa=emp, tercero=tercero)
                creados += 1
            except Exception as e:
                errores.append(
                    f'Fila {idx + 2}: {nombre} ({num_doc}) - {str(e)}'
                )

        return Response({
            'success': True,
            'empresa': emp.razon_social,
            'creados': creados,
            'vinculados': vinculados,
            'omitidos': omitidos,
            'errores': errores,
        })

class CompartirTerceroView(views.APIView):
    """POST /api/terceros/compartir/  Body: {"tercero_id": int, "empresas": [id1,...], "todas": bool}"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        from empresas.models import Empresa

        tercero_id = request.data.get("tercero_id")
        empresas_ids = request.data.get("empresas", [])
        todas = request.data.get("todas", False)

        if not tercero_id:
            return Response({"error": "tercero_id requerido"}, status=400)

        try:
            tercero = Tercero.objects.get(pk=tercero_id)
        except Tercero.DoesNotExist:
            return Response({"error": "Tercero no encontrado"}, status=404)

        if todas:
            tercero.es_compartido = True
            tercero.save(update_fields=["es_compartido"])
            # Create EmpresaTercero for ALL existing empresas
            all_emps = Empresa.objects.all()
            created = 0
            for emp in all_emps:
                _, created_flag = EmpresaTercero.objects.get_or_create(empresa=emp, tercero=tercero)
                created += int(created_flag)
            return Response({
                "status": "ok",
                "es_compartido": True,
                "empresas_vinculadas": all_emps.count(),
                "nuevas_vinculaciones": created,
            })
        else:
            if not empresas_ids:
                return Response({"error": "empresas requerido si todas=false"}, status=400)
            tercero.es_compartido = False
            tercero.save(update_fields=["es_compartido"])
            created = 0
            for eid in empresas_ids:
                try:
                    emp = Empresa.objects.get(pk=eid)
                    _, created_flag = EmpresaTercero.objects.get_or_create(empresa=emp, tercero=tercero)
                    created += int(created_flag)
                except Empresa.DoesNotExist:
                    pass
            return Response({
                "status": "ok",
                "es_compartido": False,
                "empresas_vinculadas": len(empresas_ids),
                "nuevas_vinculaciones": created,
            })
