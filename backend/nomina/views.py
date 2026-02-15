# 🎩 Don Peppini - Views Nómina
from decimal import Decimal
from datetime import date
from django.db import transaction
from django.db.models import Sum, Count, Avg
from rest_framework import viewsets, status, filters
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView
from .models import ParametrosNomina, Empleado, Nomina, LiquidacionEmpleado, DetalleHorasExtras, LiquidacionContrato
from empresas.models import Empresa
from .serializers import (
    ParametrosNominaSerializer, EmpleadoSerializer,
    NominaSerializer, LiquidacionEmpleadoSerializer,
    DetalleHorasExtrasSerializer, LiquidacionContratoSerializer,
)


from .contabilizacion import contabilizar_nomina


class ParametrosNominaViewSet(viewsets.ModelViewSet):
    """CRUD de parámetros legales de nómina por año."""
    permission_classes = [IsAuthenticated]
    queryset = ParametrosNomina.objects.all()
    serializer_class = ParametrosNominaSerializer
    pagination_class = None

    @action(detail=False, methods=['get'])
    def vigente(self, request):
        """GET /api/nomina/parametros/vigente/ — Retorna parámetros del año actual."""
        anio = int(request.query_params.get('anio', date.today().year))
        params = ParametrosNomina.del_anio(anio)
        return Response(ParametrosNominaSerializer(params).data)


class EmpleadoViewSet(viewsets.ModelViewSet):
    """CRUD de empleados por empresa."""
    permission_classes = [IsAuthenticated]
    queryset = Empleado.objects.select_related('tercero', 'empresa').all()
    serializer_class = EmpleadoSerializer
    filter_backends = [filters.SearchFilter]
    search_fields = ['tercero__nombre_razon_social', 'tercero__numero_documento', 'cargo']
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        emp = self.request.query_params.get('empresa')
        if emp:
            qs = qs.filter(empresa_id=emp)
        activo = self.request.query_params.get('activo')
        if activo is not None:
            qs = qs.filter(activo=activo.lower() == 'true')
        return qs

    def perform_create(self, serializer):
        emp = self.request.data.get('empresa')
        serializer.save(empresa_id=emp)


class NominaViewSet(viewsets.ModelViewSet):
    """Gestión de nóminas por empresa."""
    permission_classes = [IsAuthenticated]
    queryset = Nomina.objects.prefetch_related(
        'liquidaciones', 'liquidaciones__empleado', 'liquidaciones__empleado__tercero'
    ).all()
    serializer_class = NominaSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        emp = self.request.query_params.get('empresa')
        if emp:
            qs = qs.filter(empresa_id=emp)
        anio = self.request.query_params.get('anio')
        if anio:
            qs = qs.filter(anio=anio)
        return qs

    def perform_create(self, serializer):
        emp = self.request.data.get('empresa')
        serializer.save(empresa_id=emp)

    def perform_destroy(self, instance):
        """Allow delete in any state. Clean up associated asiento contable."""
        if instance.asiento_contable:
            asiento = instance.asiento_contable
            instance.asiento_contable = None
            instance.save(update_fields=['asiento_contable'])
            asiento.delete()
        instance.delete()

    @action(detail=True, methods=['post'])
    def liquidar(self, request, pk=None):
        """
        POST /api/nomina/nominas/{id}/liquidar/
        Genera las liquidaciones para todos los empleados activos de la empresa.
        Acepta novedades opcionales en el body:
        {
          "novedades": {
            "<empleado_id>": {
              "dias_trabajados": 30,
              "horas_extras": 50000,
              "comisiones": 100000,
              "bonificaciones": 0,
              "libranzas": 0,
              "otros_descuentos": 0,
              "detalle_horas": [
                {"tipo": "HED", "cantidad_horas": 5},
                {"tipo": "HEN", "cantidad_horas": 2}
              ]
            }
          }
        }
        """
        nomina = self.get_object()
        if nomina.estado not in ('borrador', 'liquidada'):
            return Response(
                {"detail": "Solo se pueden liquidar nóminas en estado borrador o liquidada."},
                status=status.HTTP_400_BAD_REQUEST
            )

        novedades = request.data.get('novedades', {})
        empleados = Empleado.objects.filter(empresa=nomina.empresa, activo=True).select_related('tercero')

        if not empleados.exists():
            return Response(
                {"detail": "No hay empleados activos en esta empresa."},
                status=status.HTTP_400_BAD_REQUEST
            )

        with transaction.atomic():
            # Limpiar liquidaciones anteriores (reliquidar)
            nomina.liquidaciones.all().delete()

            liquidaciones = []
            for emp in empleados:
                nov = novedades.get(str(emp.id), {})

                liq = LiquidacionEmpleado(
                    nomina=nomina,
                    empleado=emp,
                    dias_trabajados=nov.get('dias_trabajados', 30),
                    horas_extras=Decimal(str(nov.get('horas_extras', 0))),
                    recargos=Decimal(str(nov.get('recargos', 0))),
                    comisiones=Decimal(str(nov.get('comisiones', 0))),
                    bonificaciones=Decimal(str(nov.get('bonificaciones', 0))),
                    otros_devengados=Decimal(str(nov.get('otros_devengados', 0))),
                    libranzas=Decimal(str(nov.get('libranzas', 0))),
                    otros_descuentos=Decimal(str(nov.get('otros_descuentos', 0))),
                )
                liq.calcular()
                liquidaciones.append(liq)

                # Detalle de horas extras si se envía
                detalle_horas = nov.get('detalle_horas', [])
                for dh in detalle_horas:
                    det = DetalleHorasExtras(
                        liquidacion=liq,
                        tipo=dh['tipo'],
                        cantidad_horas=Decimal(str(dh.get('cantidad_horas', 0))),
                    )
                    det.calcular(emp.salario_base)
                    det.save()

            # Actualizar totales de la nómina
            nomina.total_devengado = sum(l.total_devengado for l in liquidaciones)
            nomina.total_deducciones = sum(l.total_deducciones for l in liquidaciones)
            nomina.total_neto = sum(l.neto_pagar for l in liquidaciones)
            nomina.total_costo_empresa = sum(l.costo_empresa for l in liquidaciones)
            nomina.estado = 'liquidada'
            nomina.fecha_liquidacion = date.today()
            nomina.save()

            # 🎩 Generar asiento contable automático
            try:
                asiento = contabilizar_nomina(nomina)
            except Exception as e:
                # No fallar la liquidación si la contabilización falla
                import logging
                logging.getLogger('nomina').warning(f"Error contabilizando nómina {nomina.id}: {e}")

        # Recargar con liquidaciones
        nomina.refresh_from_db()
        return Response(NominaSerializer(nomina).data)

    @action(detail=True, methods=['post'])
    def pagar(self, request, pk=None):
        """Marca la nómina como pagada."""
        nomina = self.get_object()
        if nomina.estado != 'liquidada':
            return Response(
                {"detail": "Solo se pueden pagar nóminas liquidadas."},
                status=status.HTTP_400_BAD_REQUEST
            )
        nomina.estado = 'pagada'
        nomina.fecha_pago = request.data.get('fecha_pago', date.today())
        nomina.save()
        return Response(NominaSerializer(nomina).data)


class LiquidacionEmpleadoViewSet(viewsets.ModelViewSet):
    """Detalle de liquidaciones individuales."""
    permission_classes = [IsAuthenticated]
    queryset = LiquidacionEmpleado.objects.select_related(
        'empleado', 'empleado__tercero', 'nomina'
    ).prefetch_related('detalle_horas').all()
    serializer_class = LiquidacionEmpleadoSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        nomina = self.request.query_params.get('nomina')
        if nomina:
            qs = qs.filter(nomina_id=nomina)
        return qs


class SimularRetencionView(APIView):
    """
    Simulador de retención en la fuente (Art. 383 ET).
    POST con datos del empleado para obtener el desglose completo.
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        from decimal import Decimal
        from .retencion_fuente import calcular_retencion_fuente

        d = request.data

        def dec(key, default='0'):
            return Decimal(str(d.get(key, default)))

        ret = calcular_retencion_fuente(
            salario_devengado=dec('salario_devengado'),
            auxilio_transporte=dec('auxilio_transporte'),
            horas_extras=dec('horas_extras'),
            comisiones=dec('comisiones'),
            bonificaciones=dec('bonificaciones'),
            otros_devengados=dec('otros_devengados'),
            uvt=dec('uvt', '52374'),
            aporte_salud_empleado=dec('aporte_salud_empleado'),
            aporte_pension_empleado=dec('aporte_pension_empleado'),
            aporte_fsp=dec('aporte_fsp'),
            tiene_dependientes=d.get('tiene_dependientes', False),
            deduccion_vivienda=dec('deduccion_vivienda'),
            deduccion_medicina_prepagada=dec('deduccion_medicina_prepagada'),
            aportes_voluntarios_pension=dec('aportes_voluntarios_pension'),
            aportes_afc=dec('aportes_afc'),
            salario_integral=d.get('salario_integral', False),
            salario_base_mensual=dec('salario_base_mensual'),
        )

        # Convert Decimals to float for JSON
        return Response({k: float(v) if isinstance(v, Decimal) else v for k, v in ret.items()})


class LiquidacionContratoViewSet(viewsets.ModelViewSet):
    """CRUD + cálculo + contabilización de liquidaciones de contrato."""
    permission_classes = [IsAuthenticated]
    serializer_class = LiquidacionContratoSerializer
    queryset = LiquidacionContrato.objects.select_related(
        'empleado', 'empleado__tercero', 'asiento_contable'
    ).all()

    def get_queryset(self):
        qs = super().get_queryset()
        empresa = self.request.query_params.get('empresa')
        if empresa:
            qs = qs.filter(empresa_id=empresa)
        return qs

    def perform_create(self, serializer):
        from datetime import datetime
        empresa_id = self.request.data.get('empresa')
        emp_id = self.request.data.get('empleado')
        empleado = Empleado.objects.get(id=emp_id)
        fecha_retiro = serializer.validated_data.get('fecha_retiro')
        if fecha_retiro and fecha_retiro < empleado.fecha_ingreso:
            from rest_framework.exceptions import ValidationError
            raise ValidationError({
                'fecha_retiro': f'No puede ser anterior a la fecha de ingreso ({empleado.fecha_ingreso})'
            })

        liq = serializer.save(
            empresa_id=empresa_id,
            salario_base=empleado.salario_base,
            fecha_ingreso=empleado.fecha_ingreso,
            tipo_contrato=empleado.tipo_contrato,
            salario_integral=empleado.salario_integral,
            trabajo_remoto=empleado.trabajo_remoto,
        )
        # Auto-calcular
        self._calcular(liq)

    def perform_update(self, serializer):
        liq = self.get_object()
        if liq.estado != 'borrador':
            from rest_framework.exceptions import ValidationError
            raise ValidationError("Solo se puede editar en estado borrador")
        liq = serializer.save()
        self._calcular(liq)

    def perform_destroy(self, instance):
        """Allow delete in any state. Clean up associated asiento contable."""
        if instance.asiento_contable:
            asiento = instance.asiento_contable
            instance.asiento_contable = None
            instance.save(update_fields=['asiento_contable'])
            asiento.delete()
        instance.delete()

    def _calcular(self, liq):
        from .liquidacion_contrato import calcular_liquidacion_contrato

        params = ParametrosNomina.del_anio(liq.fecha_retiro.year)

        r = calcular_liquidacion_contrato(
            fecha_ingreso=liq.fecha_ingreso,
            fecha_retiro=liq.fecha_retiro,
            salario_base=liq.salario_base,
            salario_integral=liq.salario_integral,
            tipo_contrato=liq.tipo_contrato,
            motivo=liq.motivo,
            fecha_fin_contrato=liq.fecha_fin_contrato,
            dias_vacaciones_disfrutados=liq.dias_vacaciones_disfrutados,
            smlv=params.smlv,
            auxilio_transporte=params.auxilio_transporte,
            trabajo_remoto=liq.trabajo_remoto,
        )

        liq.dias_ultimo_mes = r['dias_ultimo_mes']
        liq.salario_proporcional = r['salario_proporcional']
        liq.auxilio_transporte_prop = r['auxilio_transporte_prop']
        liq.dias_vacaciones_pendientes = r['dias_vacaciones_pendientes']
        liq.vacaciones = r['vacaciones']
        liq.dias_prima = r['dias_prima']
        liq.prima_servicios = r['prima_servicios']
        liq.dias_cesantias = r['dias_cesantias']
        liq.cesantias = r['cesantias']
        liq.intereses_cesantias = r['intereses_cesantias']
        liq.indemnizacion = r['indemnizacion']
        liq.deduccion_salud = r['deduccion_salud']
        liq.deduccion_pension = r['deduccion_pension']
        liq.total_devengado = r['total_devengado']
        liq.total_deducciones = r['total_deducciones']
        liq.neto_pagar = r['neto_pagar']
        liq.save()

    @action(detail=True, methods=['post'])
    def recalcular(self, request, pk=None):
        """Recalcula la liquidación (si está en borrador)."""
        liq = self.get_object()
        if liq.estado != 'borrador':
            return Response({'error': 'Solo se puede recalcular en estado borrador'},
                            status=status.HTTP_400_BAD_REQUEST)
        self._calcular(liq)
        return Response(self.get_serializer(liq).data)

    @action(detail=True, methods=['post'])
    def liquidar(self, request, pk=None):
        """Marca como liquidada y genera asiento contable."""
        liq = self.get_object()
        if liq.estado != 'borrador':
            return Response({'error': 'Solo se puede liquidar desde borrador'},
                            status=status.HTTP_400_BAD_REQUEST)

        from .contabilizacion_liquidacion import contabilizar_liquidacion_contrato
        contabilizar_liquidacion_contrato(liq)

        liq.estado = 'liquidada'
        liq.save(update_fields=['estado'])

        # Marcar empleado como inactivo y guardar fecha retiro
        emp = liq.empleado
        emp.activo = False
        emp.fecha_retiro = liq.fecha_retiro
        emp.save(update_fields=['activo', 'fecha_retiro'])

        return Response(self.get_serializer(liq).data)

    @action(detail=True, methods=['post'])
    def pagar(self, request, pk=None):
        """Marca la liquidación como pagada."""
        liq = self.get_object()
        if liq.estado != 'liquidada':
            return Response({'error': 'Debe estar liquidada para marcar como pagada'},
                            status=status.HTTP_400_BAD_REQUEST)
        liq.estado = 'pagada'
        liq.save(update_fields=['estado'])
        return Response(self.get_serializer(liq).data)


# ============================================================
# PDF COMPROBANTES
# ============================================================
from django.http import HttpResponse


class ComprobantePDFView(APIView):
    """Genera PDF de comprobante de nómina para una liquidación."""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        liq = LiquidacionEmpleado.objects.select_related(
            'empleado', 'empleado__tercero', 'nomina', 'nomina__empresa'
        ).get(pk=pk)
        empresa = liq.nomina.empresa

        from .pdf_comprobante import generar_comprobante_nomina
        buf = generar_comprobante_nomina(liq, empresa)

        nombre = f"comprobante_{liq.empleado.tercero.numero_documento}_{liq.nomina.mes}_{liq.nomina.anio}.pdf"
        response = HttpResponse(buf.read(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{nombre}"'
        return response


class ComprobanteContratoView(APIView):
    """Genera PDF de liquidación de contrato."""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        liq = LiquidacionContrato.objects.select_related(
            'empleado', 'empleado__tercero', 'empresa'
        ).get(pk=pk)

        from .pdf_comprobante import generar_comprobante_liquidacion_contrato
        buf = generar_comprobante_liquidacion_contrato(liq, liq.empresa)

        nombre = f"liquidacion_contrato_{liq.empleado.tercero.numero_documento}.pdf"
        response = HttpResponse(buf.read(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{nombre}"'
        return response


# ============================================================
# VACACIONES
# ============================================================
from .models import Vacaciones, PrimaSemestral, DetallePrima, CesantiasAnuales, DetalleCesantias
from .serializers import (
    VacacionesSerializer, PrimaSemestralSerializer,
    CesantiasAnualesSerializer,
)


class VacacionesViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = VacacionesSerializer
    queryset = Vacaciones.objects.select_related('empleado', 'empleado__tercero').all()

    def get_queryset(self):
        qs = super().get_queryset()
        empresa = self.request.query_params.get('empresa')
        if empresa:
            qs = qs.filter(empresa_id=empresa)
        empleado = self.request.query_params.get('empleado')
        if empleado:
            qs = qs.filter(empleado_id=empleado)
        return qs

    def perform_create(self, serializer):
        serializer.save(empresa_id=self.request.data.get('empresa'))

    @action(detail=False, methods=['get'])
    def saldos(self, request):
        """Calcula saldo de vacaciones de todos los empleados activos."""
        empresa = request.query_params.get('empresa')
        if not empresa:
            return Response([])

        empleados = Empleado.objects.filter(empresa_id=empresa, activo=True).select_related('tercero')
        hoy = date.today()
        resultados = []

        for emp in empleados:
            dias_trabajados = (hoy - emp.fecha_ingreso).days
            # 15 días hábiles por año (base calendario)
            dias_causados = round(dias_trabajados * 15 / 365, 1)
            dias_tomados = float(
                Vacaciones.objects.filter(
                    empleado=emp, estado__in=['aprobada', 'disfrutada']
                ).aggregate(total=Sum('dias_habiles'))['total'] or 0
            )
            pendientes = round(dias_causados - dias_tomados, 1)

            resultados.append({
                'empleado_id': emp.id,
                'empleado_nombre': emp.tercero.nombre_razon_social,
                'cargo': emp.cargo,
                'fecha_ingreso': emp.fecha_ingreso,
                'dias_causados': dias_causados,
                'dias_tomados': dias_tomados,
                'dias_pendientes': pendientes,
            })

        return Response(resultados)


# ============================================================
# PRIMA SEMESTRAL
# ============================================================
class PrimaSemestralViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = PrimaSemestralSerializer
    queryset = PrimaSemestral.objects.prefetch_related(
        'detalles', 'detalles__empleado', 'detalles__empleado__tercero'
    ).all()

    def get_queryset(self):
        qs = super().get_queryset()
        empresa = self.request.query_params.get('empresa')
        if empresa:
            qs = qs.filter(empresa_id=empresa)
        return qs

    def perform_create(self, serializer):
        empresa_id = self.request.data.get('empresa')
        anio = int(self.request.data.get('anio', date.today().year))
        semestre = int(self.request.data.get('semestre', 1))

        prima = serializer.save(empresa_id=empresa_id, anio=anio, semestre=semestre)
        self._calcular(prima)

    def _calcular(self, prima):
        from .liquidacion_contrato import dias_360_inclusive

        params = ParametrosNomina.del_anio(prima.anio)
        empleados = Empleado.objects.filter(empresa=prima.empresa, activo=True, salario_integral=False)

        if prima.semestre == 1:
            inicio_sem = date(prima.anio, 1, 1)
            fin_sem = date(prima.anio, 6, 30)
        else:
            inicio_sem = date(prima.anio, 7, 1)
            fin_sem = date(prima.anio, 12, 31)

        DetallePrima.objects.filter(prima=prima).delete()
        total = Decimal('0')

        for emp in empleados:
            fecha_inicio = max(emp.fecha_ingreso, inicio_sem)
            if fecha_inicio > fin_sem:
                continue
            dias = dias_360_inclusive(fecha_inicio, min(date.today(), fin_sem))
            dias = max(dias, 0)

            tiene_aux = emp.tiene_auxilio_transporte
            aux = params.auxilio_transporte if tiene_aux else Decimal('0')
            base = emp.salario_base + aux
            valor = (base * Decimal(str(dias)) / Decimal('360')).quantize(Decimal('0.01'))

            DetallePrima.objects.create(
                prima=prima, empleado=emp,
                salario_base=emp.salario_base,
                auxilio_transporte=aux,
                dias_trabajados=dias,
                valor_prima=valor,
            )
            total += valor

        prima.total = total
        prima.save(update_fields=['total'])

    @action(detail=True, methods=['post'])
    def liquidar(self, request, pk=None):
        prima = self.get_object()
        if prima.estado != 'borrador':
            return Response({'error': 'Solo borrador'}, status=400)

        from .contabilizacion_liquidacion import _asegurar_cuenta
        from contabilidad.models import AsientoContable, MovimientoContable
        from django.db import transaction as db_tx

        with db_tx.atomic():
            tercero = prima.empresa.terceros.first()
            concepto = f"Prima S{prima.semestre} {prima.anio}"
            asiento = AsientoContable.objects.create(
                empresa=prima.empresa, fecha=date.today(),
                tercero=tercero, concepto=concepto,
                descripcion=f"Contabilización automática — {concepto}",
            )
            gasto_cuenta = _asegurar_cuenta(prima.empresa, '510536', 'Prima de servicios')
            pasivo_cuenta = _asegurar_cuenta(prima.empresa, '262005', 'Prima de servicios por pagar')

            MovimientoContable.objects.bulk_create([
                MovimientoContable(asiento=asiento, cuenta=gasto_cuenta, debito=prima.total, credito=Decimal('0')),
                MovimientoContable(asiento=asiento, cuenta=pasivo_cuenta, debito=Decimal('0'), credito=prima.total),
            ])

            prima.asiento_contable = asiento
            prima.estado = 'liquidada'
            prima.save(update_fields=['estado', 'asiento_contable'])

        return Response(self.get_serializer(prima).data)


# ============================================================
# CESANTÍAS ANUALES
# ============================================================
class CesantiasAnualesViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = CesantiasAnualesSerializer
    queryset = CesantiasAnuales.objects.prefetch_related(
        'detalles', 'detalles__empleado', 'detalles__empleado__tercero'
    ).all()

    def get_queryset(self):
        qs = super().get_queryset()
        empresa = self.request.query_params.get('empresa')
        if empresa:
            qs = qs.filter(empresa_id=empresa)
        return qs

    def perform_create(self, serializer):
        empresa_id = self.request.data.get('empresa')
        anio = int(self.request.data.get('anio', date.today().year))

        ces = serializer.save(empresa_id=empresa_id, anio=anio)
        self._calcular(ces)

    def _calcular(self, ces):
        from .liquidacion_contrato import dias_360_inclusive

        params = ParametrosNomina.del_anio(ces.anio)
        empleados = Empleado.objects.filter(empresa=ces.empresa, activo=True, salario_integral=False)
        inicio = date(ces.anio, 1, 1)
        fin = date(ces.anio, 12, 31)

        DetalleCesantias.objects.filter(cesantias=ces).delete()
        total_c = Decimal('0')
        total_i = Decimal('0')

        for emp in empleados:
            fecha_inicio = max(emp.fecha_ingreso, inicio)
            if fecha_inicio > fin:
                continue
            dias = dias_360_inclusive(fecha_inicio, min(date.today(), fin))
            dias = max(dias, 0)

            tiene_aux = emp.tiene_auxilio_transporte
            aux = params.auxilio_transporte if tiene_aux else Decimal('0')
            base = emp.salario_base + aux

            cesantia = (base * Decimal(str(dias)) / Decimal('360')).quantize(Decimal('0.01'))
            intereses = (cesantia * Decimal(str(dias)) * Decimal('0.12') / Decimal('360')).quantize(Decimal('0.01'))

            DetalleCesantias.objects.create(
                cesantias=ces, empleado=emp,
                salario_base=emp.salario_base,
                auxilio_transporte=aux,
                dias_trabajados=dias,
                valor_cesantias=cesantia,
                valor_intereses=intereses,
            )
            total_c += cesantia
            total_i += intereses

        ces.total_cesantias = total_c
        ces.total_intereses = total_i
        ces.save(update_fields=['total_cesantias', 'total_intereses'])

    @action(detail=True, methods=['post'])
    def liquidar(self, request, pk=None):
        ces = self.get_object()
        if ces.estado != 'borrador':
            return Response({'error': 'Solo borrador'}, status=400)

        from .contabilizacion_liquidacion import _asegurar_cuenta
        from contabilidad.models import AsientoContable, MovimientoContable
        from django.db import transaction as db_tx

        with db_tx.atomic():
            tercero = ces.empresa.terceros.first()
            concepto = f"Cesantías + Intereses {ces.anio}"
            asiento = AsientoContable.objects.create(
                empresa=ces.empresa, fecha=date.today(),
                tercero=tercero, concepto=concepto,
                descripcion=f"Contabilización automática — {concepto}",
            )
            gasto_ces = _asegurar_cuenta(ces.empresa, '510530', 'Cesantías')
            gasto_int = _asegurar_cuenta(ces.empresa, '510533', 'Intereses sobre cesantías')
            pasivo_ces = _asegurar_cuenta(ces.empresa, '261005', 'Cesantías consolidadas')
            pasivo_int = _asegurar_cuenta(ces.empresa, '261505', 'Intereses sobre cesantías')

            movs = []
            if ces.total_cesantias > 0:
                movs.append(MovimientoContable(asiento=asiento, cuenta=gasto_ces, debito=ces.total_cesantias, credito=Decimal('0')))
                movs.append(MovimientoContable(asiento=asiento, cuenta=pasivo_ces, debito=Decimal('0'), credito=ces.total_cesantias))
            if ces.total_intereses > 0:
                movs.append(MovimientoContable(asiento=asiento, cuenta=gasto_int, debito=ces.total_intereses, credito=Decimal('0')))
                movs.append(MovimientoContable(asiento=asiento, cuenta=pasivo_int, debito=Decimal('0'), credito=ces.total_intereses))

            if movs:
                MovimientoContable.objects.bulk_create(movs)

            ces.asiento_contable = asiento
            ces.estado = 'liquidada'
            ces.save(update_fields=['estado', 'asiento_contable'])

        return Response(self.get_serializer(ces).data)


# ============================================================
# DASHBOARD NÓMINA (Resumen)
# ============================================================
class DashboardNominaView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        empresa = request.query_params.get('empresa')
        if not empresa:
            return Response({})

        from django.db.models import Sum, Count, Q, Avg

        empleados = Empleado.objects.filter(empresa_id=empresa)
        activos = empleados.filter(activo=True)

        # Última nómina
        ultima_nomina = Nomina.objects.filter(empresa_id=empresa).order_by('-anio', '-mes').first()

        # Totales última nómina
        ultima_data = {}
        if ultima_nomina:
            agg = ultima_nomina.liquidaciones.aggregate(
                total_dev=Sum('total_devengado'),
                total_ded=Sum('total_deducciones'),
                total_neto=Sum('neto_pagar'),
                total_costo=Sum('costo_empresa'),
                count=Count('id'),
            )
            ultima_data = {
                'periodo': f"{ultima_nomina.mes}/{ultima_nomina.anio}",
                'estado': ultima_nomina.get_estado_display(),
                **agg,
            }

        # Saldo vacaciones promedio
        hoy = date.today()
        vac_data = []
        for emp in activos:
            dias_trab = (hoy - emp.fecha_ingreso).days
            causados = dias_trab * 15 / 365
            tomados = float(Vacaciones.objects.filter(
                empleado=emp, estado__in=['aprobada', 'disfrutada']
            ).aggregate(t=Sum('dias_habiles'))['t'] or 0)
            vac_data.append(causados - tomados)

        return Response({
            'empleados_activos': activos.count(),
            'empleados_inactivos': empleados.filter(activo=False).count(),
            'salario_promedio': float(activos.aggregate(avg=Avg('salario_base'))['avg'] or 0),
            'nomina_total': float(activos.aggregate(total=Sum('salario_base'))['total'] or 0),
            'ultima_nomina': ultima_data,
            'vacaciones_pendientes_prom': round(sum(vac_data) / len(vac_data), 1) if vac_data else 0,
            'liquidaciones_contrato': LiquidacionContrato.objects.filter(
                empresa_id=empresa, estado='borrador').count(),
        })


# ============================================================
# IMPORTAR EMPLEADOS DESDE EXCEL
# ============================================================
class ImportarEmpleadosView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        import openpyxl
        from terceros.models import Tercero

        empresa_id = request.data.get('empresa')
        archivo = request.FILES.get('archivo')
        if not archivo:
            return Response({'error': 'No se envió archivo'}, status=400)

        try:
            wb = openpyxl.load_workbook(archivo)
            ws = wb.active
            headers = [str(c.value or '').strip().lower() for c in ws[1]]

            creados = 0
            errores = []

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                data = dict(zip(headers, row))
                try:
                    documento = str(data.get('documento', '') or '').strip()
                    nombre = str(data.get('nombre', '') or '').strip()
                    if not documento or not nombre:
                        continue

                    tipo_doc = str(data.get('tipo_documento', 'CC') or 'CC').strip().upper()
                    tercero, _ = Tercero.objects.get_or_create(
                        empresa_id=empresa_id,
                        numero_documento=documento,
                        defaults={
                            'nombre_razon_social': nombre,
                            'tipo_documento': tipo_doc,
                            'tipo_tercero': 'empleado',
                        }
                    )

                    salario = Decimal(str(data.get('salario', 0) or 0))
                    fecha_ing = data.get('fecha_ingreso')
                    if isinstance(fecha_ing, str):
                        from datetime import datetime
                        fecha_ing = datetime.strptime(fecha_ing, '%Y-%m-%d').date()

                    emp, created = Empleado.objects.get_or_create(
                        empresa_id=empresa_id,
                        tercero=tercero,
                        defaults={
                            'salario_base': salario,
                            'fecha_ingreso': fecha_ing or date.today(),
                            'tipo_contrato': str(data.get('tipo_contrato', 'IND') or 'IND').strip().upper()[:3],
                            'cargo': str(data.get('cargo', '') or '').strip(),
                            'eps': str(data.get('eps', '') or '').strip(),
                            'afp': str(data.get('afp', '') or '').strip(),
                        }
                    )
                    if created:
                        creados += 1
                except Exception as e:
                    errores.append(f"Fila {row_idx}: {str(e)}")

            return Response({
                'creados': creados,
                'errores': errores,
                'mensaje': f"Se importaron {creados} empleados" + (f" con {len(errores)} errores" if errores else ""),
            })
        except Exception as e:
            return Response({'error': str(e)}, status=400)


# ============================================================
# 🎩 PILA — Seguridad Social y Parafiscales
# ============================================================

class PILAPreviewView(APIView):
    """Preview de PILA: calcula montos sin crear asiento."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        empresa_id = request.query_params.get('empresa')
        anio = request.query_params.get('anio')
        mes = request.query_params.get('mes')

        if not all([empresa_id, anio, mes]):
            return Response({'error': 'empresa, anio y mes son requeridos'}, status=400)

        try:
            empresa = Empresa.objects.get(id=empresa_id)
        except Empresa.DoesNotExist:
            return Response({'error': 'Empresa no encontrada'}, status=404)

        from .pila import calcular_pila
        result = calcular_pila(empresa, int(anio), int(mes))

        if 'error' in result:
            return Response(result, status=400)

        return Response(result)


class PILACausarView(APIView):
    """Info: la causación ya se hace al liquidar nómina."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response({
            'info': 'La causación de aportes se genera automáticamente al liquidar la nómina. '
                    'Use el endpoint de pago para registrar el desembolso.'
        })


class PILAPagarView(APIView):
    """Genera asiento de pago de PILA."""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        empresa_id = request.data.get('empresa')
        anio = request.data.get('anio')
        mes = request.data.get('mes')
        fecha = request.data.get('fecha')
        cuenta_banco = request.data.get('cuenta_banco')

        if not all([empresa_id, anio, mes, fecha, cuenta_banco]):
            return Response({'error': 'empresa, anio, mes, fecha y cuenta_banco son requeridos'}, status=400)

        try:
            empresa = Empresa.objects.get(id=empresa_id)
        except Empresa.DoesNotExist:
            return Response({'error': 'Empresa no encontrada'}, status=404)

        from datetime import datetime
        try:
            fecha_obj = datetime.strptime(fecha, '%Y-%m-%d').date()
        except ValueError:
            return Response({'error': 'Fecha inválida (YYYY-MM-DD)'}, status=400)

        from .pila import pagar_pila
        result = pagar_pila(empresa, int(anio), int(mes), fecha_obj, cuenta_banco, request.user)

        if 'error' in result:
            return Response(result, status=400)

        return Response(result, status=201)
